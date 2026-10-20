#!/usr/bin/env python3
"""Generate the yearly habit tracker spreadsheet (habit-tracker.xlsx).

    python3 tools/make_xlsx.py [output.xlsx] [--year 2026]

One Setup sheet (year + up to 10 habits typed once), an Overview sheet and one sheet per month.
Everything is driven by formulas, so changing the year on Setup re-dates every month.
Only functions that work in Excel 2010+, LibreOffice and Google Sheets are used
(DATE, DAY, MONTH, YEAR, TODAY, WEEKDAY, INDEX, CHOOSE, SUMPRODUCT, MIN, MAX, ROUND, AVERAGE, IFERROR).
Requires openpyxl (pip install openpyxl).
"""
import argparse
import os

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

# hmsk.life Letterpress A1 "Ink Blue" tokens (public/letterpress.css), adapted for a white sheet.
# Header text is white on the accent (8.6:1), ticks and today are ink on mark (10:1), small text ink-2 on white (7.7:1).
INK = "151B24"
INK_2 = "4A5562"
ACCENT = "1E4C8A"
MARK = "A8C8F2"
TINT = "DCE6F3"
RULE = "CBD2D9"
WEEKEND = "EFF1F2"
OUTSIDE = "E2E6E9"

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
N_HABITS = 10
FIRST_DAY_COL = 2           # B
LAST_DAY_COL = 32           # AF (31 days)
COL_DONE, COL_GOAL, COL_PCT, COL_STREAK, COL_END = 33, 34, 35, 36, 37   # AG AH AI AJ AK
ROW_DATE, ROW_WD, ROW_H1 = 3, 4, 5
ROW_HN = ROW_H1 + N_HABITS - 1      # 14
ROW_DAYPCT = ROW_HN + 1             # 15
ROW_HELP0 = ROW_DAYPCT + 3          # 18: hidden running-streak helper rows 18..27
SETUP_YEAR = "Setup!$C$3"
SETUP_FIRST = 6                     # habits on Setup!B6:B15, days per week in C6:C15
SETUP_COUNT = "Setup!$C$17"

L = get_column_letter
DAY_FIRST, DAY_LAST = L(FIRST_DAY_COL), L(LAST_DAY_COL)

dark = PatternFill("solid", fgColor=ACCENT)
input_fill = PatternFill("solid", fgColor=TINT)
thin = Side(style="thin", color=RULE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")
F_TITLE = Font(name="Calibri", size=18, bold=True, color=INK)
F_HEAD = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_WD = Font(name="Calibri", size=9, color=MARK)  # on the accent header: 5.1:1
F_SMALL = Font(name="Calibri", size=9, color=INK_2)
F_BOLD = Font(name="Calibri", size=11, bold=True)
F_TICK = Font(name="Calibri", size=11, bold=True, color=INK)

BLUE_SCALE = dict(start_type="num", start_value=0, start_color="F8F9FA",
                  mid_type="num", mid_value=0.5, mid_color=MARK,
                  end_type="num", end_value=1, end_color="7FA6DC")  # ink text stays ≥ 6.9:1 on every step


def page_setup(ws, area):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_area = area
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True


def build_setup(wb, year, examples):
    ws = wb.active
    ws.title = "Setup"
    ws.sheet_properties.tabColor = MARK
    ws["A1"] = "Habit Tracker"
    ws["A1"].font = F_TITLE
    ws["D1"] = "hmsk.life/tools/habit-tracker"
    ws["D1"].font = F_SMALL
    ws["A3"] = "Year / Jahr"
    ws["A3"].font = F_BOLD
    ws["C3"] = year
    ws["C3"].font = Font(name="Calibri", size=14, bold=True)
    ws["C3"].fill = PatternFill("solid", fgColor=MARK)
    ws["C3"].alignment = center
    ws["C3"].border = box
    ws["D3"] = "Change the year here, every month sheet follows. / Jahr hier ändern, alle Monate folgen."
    ws["D3"].font = F_SMALL
    dv_year = DataValidation(type="whole", operator="between", formula1="1900", formula2="9999", allow_blank=False,
                             showErrorMessage=True, errorTitle="Year", error="Enter a year, e.g. 2027.")
    ws.add_data_validation(dv_year)
    dv_year.add("C3")

    for col, text in ((1, "#"), (2, "Habit / Gewohnheit"), (3, "Days per week / Tage pro Woche (1-7)")):
        c = ws.cell(5, col, text)
        c.font = F_HEAD
        c.fill = dark
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[5].height = 30
    dv_days = DataValidation(type="whole", operator="between", formula1="1", formula2="7", allow_blank=True,
                             showErrorMessage=True, errorTitle="Days per week", error="1 to 7 (7 = every day).")
    ws.add_data_validation(dv_days)
    for i in range(N_HABITS):
        r = SETUP_FIRST + i
        ws.cell(r, 1, i + 1).font = F_SMALL
        name, days = examples[i] if i < len(examples) else (None, 7)
        b = ws.cell(r, 2, name)
        c = ws.cell(r, 3, days)
        for cell in (b, c):
            cell.fill = input_fill
            cell.border = box
        c.alignment = center
        dv_days.add(c.coordinate)
        ws.row_dimensions[r].height = 20
    ws["A17"] = "Habits set up"
    ws["A17"].font = F_SMALL
    ws["C17"] = f"=SUMPRODUCT((B{SETUP_FIRST}:B{SETUP_FIRST + N_HABITS - 1}<>\"\")*1)"
    ws["C17"].alignment = center
    ws["C17"].font = F_SMALL

    help_lines = [
        "How to use",
        "1. Type your habits above (up to 10). Replace or delete the three examples. Days per week: 7 = daily, 3 = three times a week.",
        "2. On each month sheet, put x (or pick ✓ from the list) in a cell when you did the habit that day. Empty = not done.",
        "3. The month sheet shows done days, the goal for the month, % of the goal and the current streak (days in a row).",
        "4. Overview: % per month, the whole year, current and best streak per habit.",
        "",
        "So geht's",
        "1. Oben die Gewohnheiten eintragen (bis zu 10). Die drei Beispiele ersetzen oder löschen. Tage pro Woche: 7 = täglich.",
        "2. Auf dem Monatsblatt ein x (oder ✓ aus der Liste) in die Zelle setzen, wenn Sie die Gewohnheit an dem Tag erledigt haben.",
        "3. Das Monatsblatt zeigt erledigte Tage, das Monatsziel, % vom Ziel und die aktuelle Serie (Tage in Folge).",
        "4. Overview: % pro Monat, das ganze Jahr, aktuelle und beste Serie je Gewohnheit.",
        "",
        "Prefer the browser? hmsk.life/apps/habit-tracker/ (free, no signup, data stays on your device).",
    ]
    for i, line in enumerate(help_lines):
        c = ws.cell(19 + i, 2, line)
        c.font = F_BOLD if line in ("How to use", "So geht's") else Font(name="Calibri", size=10)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 60
    ws.freeze_panes = "A6"
    page_setup(ws, f"A1:D{19 + len(help_lines)}")


def build_month(wb, m):
    ws = wb.create_sheet(MONTHS[m - 1])
    prev = MONTHS[m - 2] if m > 1 else None
    ws.sheet_properties.tabColor = INK_2

    ws["A1"] = f"=DATE({SETUP_YEAR},{m},1)"
    ws["A1"].number_format = "mmmm yyyy"
    ws["A1"].font = F_TITLE
    ws["A1"].alignment = Alignment(horizontal="left")
    ws.merge_cells("A1:L1")
    ws["N1"] = "x or ✓ = done. Empty = not done. / x oder ✓ = erledigt."
    ws["N1"].font = F_SMALL

    # header rows
    for col, text in ((1, "Habit"), (COL_DONE, "Done"), (COL_GOAL, "Goal"), (COL_PCT, "%"), (COL_STREAK, "Streak")):
        c = ws.cell(ROW_DATE, col, text)
        c.font = F_HEAD
        c.fill = dark
        c.alignment = center if col > 1 else Alignment(vertical="center")
        ws.cell(ROW_WD, col).fill = dark
    ws.cell(ROW_WD, COL_STREAK, "days").font = F_WD
    ws.cell(ROW_WD, COL_STREAK).alignment = center
    for col in range(FIRST_DAY_COL, LAST_DAY_COL + 1):
        letter, before = L(col), L(col - 1)
        if col == FIRST_DAY_COL:
            f = f"=DATE({SETUP_YEAR},{m},1)"
        else:
            f = f'=IF({before}{ROW_DATE}="","",IF(MONTH({before}{ROW_DATE}+1)={m},{before}{ROW_DATE}+1,""))'
        d = ws.cell(ROW_DATE, col, f)
        d.number_format = "d"
        d.font = F_HEAD
        d.fill = dark
        d.alignment = center
        w = ws.cell(ROW_WD, col, f'=IF({letter}${ROW_DATE}="","",{letter}${ROW_DATE})')
        w.number_format = "ddd"
        w.font = F_WD
        w.fill = dark
        w.alignment = center

    days_in_month = f"DAY(DATE({SETUP_YEAR},{m + 1},0))"
    for i in range(N_HABITS):
        r = ROW_H1 + i
        s = SETUP_FIRST + i
        hr = ROW_HELP0 + i
        name = ws.cell(r, 1, f'=IF(Setup!$B{s}="","",Setup!$B{s})')
        name.font = F_BOLD
        name.alignment = Alignment(vertical="center", shrink_to_fit=True)
        ws.row_dimensions[r].height = 22
        for col in range(FIRST_DAY_COL, LAST_DAY_COL + 1):
            c = ws.cell(r, col)
            c.alignment = center
            c.border = box
            c.font = F_TICK
        rng = f"${DAY_FIRST}{r}:${DAY_LAST}{r}"
        ws.cell(r, COL_DONE, f'=IF($A{r}="","",SUMPRODUCT(({rng}<>"")*(${DAY_FIRST}${ROW_DATE}:${DAY_LAST}${ROW_DATE}<>"")))')
        per_week = f'IF(Setup!$C{s}="",7,Setup!$C{s})'
        ws.cell(r, COL_GOAL, f'=IF($A{r}="","",MAX(1,ROUND({days_in_month}*{per_week}/7,0)))')
        p = ws.cell(r, COL_PCT, f'=IF($A{r}="","",MIN(1,{L(COL_DONE)}{r}/{L(COL_GOAL)}{r}))')
        p.number_format = "0%"
        # running streak helper (hidden): days in a row up to each day, carried over from last month
        prev_end = f"{prev}!${L(COL_END)}{r}" if prev else "0"
        for col in range(FIRST_DAY_COL, LAST_DAY_COL + 1):
            letter, before = L(col), L(col - 1)
            if col == FIRST_DAY_COL:
                f = f'=IF(AND({letter}${ROW_DATE}<>"",{letter}{r}<>""),N({prev_end})+1,0)'
            else:
                f = f'=IF({letter}${ROW_DATE}="",{before}{hr},IF({letter}{r}<>"",{before}{hr}+1,0))'
            ws.cell(hr, col, f)
        ws.cell(r, COL_END, f"={DAY_LAST}{hr}")
        helper = f"${DAY_FIRST}{hr}:${DAY_LAST}{hr}"
        today_in_month = f"AND(YEAR(TODAY())={SETUP_YEAR},MONTH(TODAY())={m})"
        yesterday = f"IF(DAY(TODAY())=1,N({prev_end}),INDEX({helper},1,DAY(TODAY())-1))"
        streak = (f'=IF($A{r}="","",IF({today_in_month},'
                  f'IF(INDEX({rng},1,DAY(TODAY()))<>"",INDEX({helper},1,DAY(TODAY())),{yesterday}),'
                  f'IF(TODAY()>DATE({SETUP_YEAR},{m + 1},0),${L(COL_END)}{r},"")))')
        ws.cell(r, COL_STREAK, streak)
        for col in (COL_DONE, COL_GOAL, COL_PCT, COL_STREAK):
            ws.cell(r, col).alignment = center
            ws.cell(r, col).border = box

    # daily % row
    ws.cell(ROW_DAYPCT, 1, "Done that day").font = F_SMALL
    for col in range(FIRST_DAY_COL, LAST_DAY_COL + 1):
        letter = L(col)
        f = (f'=IF(OR({letter}${ROW_DATE}="",{SETUP_COUNT}=0),"",IF({letter}${ROW_DATE}>TODAY(),"",'
             f'SUMPRODUCT(({letter}${ROW_H1}:{letter}${ROW_HN}<>"")*($A${ROW_H1}:$A${ROW_HN}<>""))/{SETUP_COUNT}))')
        c = ws.cell(ROW_DAYPCT, col, f)
        c.number_format = "0%"
        c.font = Font(name="Calibri", size=8, color=INK_2)
        c.alignment = center
    month_pct = ws.cell(ROW_DAYPCT, COL_PCT, f'=IFERROR(AVERAGE({L(COL_PCT)}{ROW_H1}:{L(COL_PCT)}{ROW_HN}),"")')
    month_pct.number_format = "0%"
    month_pct.font = F_BOLD
    month_pct.alignment = center
    ws.cell(ROW_DAYPCT, COL_GOAL, "all").font = F_SMALL
    ws.cell(ROW_DAYPCT, COL_GOAL).alignment = center

    ws.cell(ROW_HELP0 - 1, 1, "helper: days in a row (do not edit)").font = F_SMALL
    for r in range(ROW_HELP0 - 1, ROW_HELP0 + N_HABITS):
        ws.row_dimensions[r].hidden = True
    ws.column_dimensions[L(COL_END)].hidden = True

    # ticks: dropdown with x / ✓ / 1, any text still counts as done
    dv = DataValidation(type="list", formula1='"x,✓,1"', allow_blank=True, showErrorMessage=True, errorStyle="warning",
                        errorTitle="Tick", error="Use x or ✓ to mark a day as done.")
    ws.add_data_validation(dv)
    dv.add(f"{DAY_FIRST}{ROW_H1}:{DAY_LAST}{ROW_HN}")

    # conditional formatting (first match wins)
    grid = f"{DAY_FIRST}{ROW_DATE}:{DAY_LAST}{ROW_DAYPCT}"
    ticks = f"{DAY_FIRST}{ROW_H1}:{DAY_LAST}{ROW_HN}"
    ws.conditional_formatting.add(grid, FormulaRule(formula=[f'{DAY_FIRST}${ROW_DATE}=""'], stopIfTrue=True,
                                                    fill=PatternFill("solid", fgColor=OUTSIDE, bgColor=OUTSIDE)))
    ws.conditional_formatting.add(ticks, FormulaRule(formula=[f'{DAY_FIRST}{ROW_H1}<>""'], stopIfTrue=True,
                                                     fill=PatternFill("solid", fgColor=MARK, bgColor=MARK), font=Font(bold=True, color=INK)))
    ws.conditional_formatting.add(f"{DAY_FIRST}{ROW_DATE}:{DAY_LAST}{ROW_WD}",
                                  FormulaRule(formula=[f"{DAY_FIRST}${ROW_DATE}=TODAY()"], stopIfTrue=True,
                                              fill=PatternFill("solid", fgColor=MARK, bgColor=MARK), font=Font(bold=True, color=INK)))
    ws.conditional_formatting.add(ticks, FormulaRule(formula=[f"WEEKDAY({DAY_FIRST}${ROW_DATE},2)>5"],
                                                     fill=PatternFill("solid", fgColor=WEEKEND, bgColor=WEEKEND)))
    ws.conditional_formatting.add(f"{L(COL_PCT)}{ROW_H1}:{L(COL_PCT)}{ROW_DAYPCT}", ColorScaleRule(**BLUE_SCALE))
    ws.conditional_formatting.add(f"{DAY_FIRST}{ROW_DAYPCT}:{DAY_LAST}{ROW_DAYPCT}", ColorScaleRule(**BLUE_SCALE))

    ws.column_dimensions["A"].width = 26
    for col in range(FIRST_DAY_COL, LAST_DAY_COL + 1):
        ws.column_dimensions[L(col)].width = 4.2
    for col in (COL_DONE, COL_GOAL, COL_PCT, COL_STREAK):
        ws.column_dimensions[L(col)].width = 7.5
    ws.row_dimensions[ROW_DATE].height = 18
    ws.freeze_panes = f"{DAY_FIRST}{ROW_H1}"
    page_setup(ws, f"A1:{L(COL_STREAK)}{ROW_DAYPCT}")
    return ws


def build_overview(wb):
    ws = wb.create_sheet("Overview", 1)
    ws.sheet_properties.tabColor = ACCENT
    ws["A1"] = "Habit Tracker"
    ws["A1"].font = F_TITLE
    ws["C1"] = f"={SETUP_YEAR}"
    ws["C1"].font = F_TITLE
    ws["C1"].alignment = Alignment(horizontal="left")
    heads = ["Habit"] + [None] * 12 + ["Done", "Year %", "Streak now", "Best streak"]
    for col, text in enumerate(heads, start=1):
        c = ws.cell(3, col, text)
        if 2 <= col <= 13:
            c.value = f"=DATE({SETUP_YEAR},{col - 1},1)"
            c.number_format = "mmm"
        c.font = F_HEAD
        c.fill = dark
        c.alignment = center if col > 1 else Alignment(vertical="center")
    ws.row_dimensions[3].height = 20
    for i in range(N_HABITS):
        r = 4 + i
        mr = ROW_H1 + i
        s = SETUP_FIRST + i
        hr = ROW_HELP0 + i
        ws.cell(r, 1, f'=IF(Setup!$B{s}="","",Setup!$B{s})').font = F_BOLD
        for m in range(1, 13):
            c = ws.cell(r, 1 + m, f'=IF(OR($A{r}="",DATE({SETUP_YEAR},{m},1)>TODAY()),"",{MONTHS[m - 1]}!${L(COL_PCT)}{mr})')
            c.number_format = "0%"
            c.alignment = center
            c.border = box
        ws.cell(r, 14, f'=IF($A{r}="","",' + "+".join(f"{mo}!${L(COL_DONE)}{mr}" for mo in MONTHS) + ")")
        elapsed = f"MAX(0,MIN(TODAY(),DATE({SETUP_YEAR},12,31))-DATE({SETUP_YEAR},1,1)+1)"
        per_week = f'IF(Setup!$C{s}="",7,Setup!$C{s})'
        yp = ws.cell(r, 15, f'=IF(OR($A{r}="",{elapsed}=0),"",MIN(1,N{r}/MAX(1,{elapsed}*{per_week}/7)))')
        yp.number_format = "0%"
        now = ("=IF($A{r}=\"\",\"\",IF(YEAR(TODAY())={y},CHOOSE(MONTH(TODAY()),{months}),IF(YEAR(TODAY())>{y},Dec!${end}{mr},0)))"
               .format(r=r, y=SETUP_YEAR, end=L(COL_END), mr=mr,
                       months=",".join(f"{mo}!${L(COL_STREAK)}{mr}" for mo in MONTHS)))
        ws.cell(r, 16, now)
        ws.cell(r, 17, f'=IF($A{r}="","",MAX(' + ",".join(f"{mo}!${DAY_FIRST}{hr}:${DAY_LAST}{hr}" for mo in MONTHS) + "))")
        for col in range(14, 18):
            ws.cell(r, col).alignment = center
            ws.cell(r, col).border = box
        ws.row_dimensions[r].height = 20
    tr = 4 + N_HABITS
    ws.cell(tr, 1, "All habits").font = F_BOLD
    for m in range(1, 13):
        col = L(1 + m)
        c = ws.cell(tr, 1 + m, f'=IFERROR(AVERAGE({col}4:{col}{tr - 1}),"")')
        c.number_format = "0%"
        c.alignment = center
        c.font = F_BOLD
    c = ws.cell(tr, 15, f'=IFERROR(AVERAGE(O4:O{tr - 1}),"")')
    c.number_format = "0%"
    c.alignment = center
    c.font = F_BOLD
    ws.cell(tr + 2, 1, "% = share of the goal (days per week on Setup). Streak = days in a row, today counts once ticked.").font = F_SMALL
    ws.conditional_formatting.add(f"B4:M{tr}", ColorScaleRule(**BLUE_SCALE))
    ws.conditional_formatting.add(f"O4:O{tr}", ColorScaleRule(**BLUE_SCALE))
    ws.column_dimensions["A"].width = 28
    for col in range(2, 14):
        ws.column_dimensions[L(col)].width = 7
    for col, w in zip("NOPQ", (8, 9, 11, 11)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B4"
    page_setup(ws, f"A1:Q{tr + 2}")


def make(path, year):
    wb = Workbook()
    examples = [("Drink 2 l of water", 7), ("Walk 30 minutes", 5), ("Read 10 pages", 7)]
    build_setup(wb, year, examples)
    for m in range(1, 13):
        build_month(wb, m)
    build_overview(wb)
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = "Habit Tracker"
    wb.properties.creator = "Hafiz Muhammad Sharoz Khalid (hmsk.life)"
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("out", nargs="?", default="habit-tracker.xlsx")
    ap.add_argument("--year", type=int, default=2026)
    a = ap.parse_args()
    print(make(a.out, a.year))
